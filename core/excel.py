from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def workbook_response(filename: str):
    buffer = BytesIO()
    return buffer, filename


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True)


def export_rows(sheet_name: str, headers: list, rows: list) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    style_header(ws)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def parse_upload(file, expected_headers: list | None = None) -> list[dict]:
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    if expected_headers:
        for exp in expected_headers:
            if exp not in headers:
                raise ValueError(f"Missing column: {exp}")
    data = []
    for row in rows[1:]:
        if not any(row):
            continue
        data.append({headers[i]: row[i] for i in range(len(headers)) if i < len(row)})
    return data
